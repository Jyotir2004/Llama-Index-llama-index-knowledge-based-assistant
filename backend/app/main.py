import os
from pathlib import Path
from uuid import uuid4

os.environ["ANONYMIZED_TELEMETRY"] = "False"

import chromadb
from dotenv import load_dotenv
from docx import Document as DocxDocument
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader

from llama_index.core import Document, StorageContext, VectorStoreIndex
from llama_index.embeddings.huggingface_api import HuggingFaceInferenceAPIEmbedding
from llama_index.llms.groq import Groq
from llama_index.vector_stores.chroma import ChromaVectorStore

BASE_DIR = Path(__file__).resolve().parent.parent.parent
UPLOAD_DIR = BASE_DIR / "uploads"
CHROMA_DIR = BASE_DIR / "chroma_db"
ALLOWED_EXTS = {".pdf", ".txt", ".docx", ".csv", ".json", ".md"}
COLLECTION_NAME = "rag_documents_minilm"
DEFAULT_EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"

load_dotenv(BASE_DIR / ".env", override=True)

app = FastAPI(title="LlamaIndex RAG App")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.state.chat_history = []
app.state.index = None
app.state.embed_model = None

UPLOAD_DIR.mkdir(exist_ok=True, parents=True)
CHROMA_DIR.mkdir(exist_ok=True, parents=True)


def get_groq_api_key() -> str:
    key = os.getenv("GROQ_API_KEY")
    if not key:
        raise RuntimeError("GROQ_API_KEY is missing. Add it to the .env file.")
    return key


def extract_text_from_file(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    try:
        if suffix in {".txt", ".md", ".json", ".csv"}:
            return file_path.read_text(encoding="utf-8", errors="ignore")

        if suffix == ".pdf":
            reader = PdfReader(str(file_path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)

        if suffix == ".docx":
            doc = DocxDocument(str(file_path))
            return "\n".join(paragraph.text for paragraph in doc.paragraphs)
    except Exception:
        return ""

    return ""


def _sync_embed_single(model, text: str) -> list[float]:
    embedding = model._sync_client.feature_extraction(text)
    if hasattr(embedding, "tolist"):
        return embedding.tolist()
    if hasattr(embedding, "__iter__"):
        return list(embedding)
    return embedding


def get_embedding_model():
    if app.state.embed_model is None:
        token = os.getenv("HF_TOKEN")
        if not token:
            raise RuntimeError(
                "HF_TOKEN is missing. Add your Hugging Face token to the local .env file "
                "for the all-MiniLM-L6-v2 embedding model."
            )
        embed_model = HuggingFaceInferenceAPIEmbedding(
            model_name=os.getenv("HF_EMBEDDING_MODEL", DEFAULT_EMBEDDING_MODEL),
            token=token,
        )
        from llama_index.embeddings.huggingface_api.base import format_query, format_text

        embed_model._get_query_embedding = lambda q: _sync_embed_single(
            embed_model, format_query(q, embed_model.model_name, embed_model.query_instruction)
        )
        embed_model._get_text_embedding = lambda t: _sync_embed_single(
            embed_model, format_text(t, embed_model.model_name, embed_model.text_instruction)
        )
        embed_model._get_text_embeddings = lambda texts: [
            embed_model._get_text_embedding(t) for t in texts
        ]
        app.state.embed_model = embed_model
    return app.state.embed_model


def get_llm():
    model_name = os.getenv("GROQ_LLM_MODEL", "openai/gpt-oss-120b")
    return Groq(
        model=model_name,
        api_key=get_groq_api_key(),
        temperature=0.1,
    )


def get_chroma_collection():
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    return client.get_or_create_collection(name=COLLECTION_NAME)


def index_uploaded_file(file_path: Path, file_name: str, file_id: str) -> VectorStoreIndex:
    text = extract_text_from_file(file_path)
    if not text.strip():
        raise ValueError("No readable text was found in this document.")

    embedding_model = get_embedding_model()
    chroma_collection = get_chroma_collection()
    vector_store = ChromaVectorStore(chroma_collection=chroma_collection)
    storage_context = StorageContext.from_defaults(vector_store=vector_store)
    document = Document(
        text=text,
        metadata={"file_name": file_name, "file_id": file_id},
    )

    try:
        VectorStoreIndex.from_documents(
            [document],
            storage_context=storage_context,
            embed_model=embedding_model,
            show_progress=False,
        )
    except Exception:
        chroma_collection.delete(where={"file_id": file_id})
        raise

    app.state.index = VectorStoreIndex.from_vector_store(
        vector_store,
        embed_model=embedding_model,
    )
    return app.state.index


def ensure_index():
    chroma_collection = get_chroma_collection()
    if chroma_collection.count() == 0:
        files = [
            path
            for path in sorted(UPLOAD_DIR.iterdir())
            if path.is_file() and path.suffix.lower() in ALLOWED_EXTS
        ]
        documents = []
        for path in files:
            text = extract_text_from_file(path)
            if text.strip():
                documents.append(
                    Document(
                        text=text,
                        metadata={"file_name": path.name, "file_id": f"existing-{path.name}"},
                    )
                )
        if not documents:
            app.state.index = None
            return None

        embedding_model = get_embedding_model()
        vector_store = ChromaVectorStore(chroma_collection=chroma_collection)
        storage_context = StorageContext.from_defaults(vector_store=vector_store)
        file_ids = [document.metadata["file_id"] for document in documents]
        try:
            VectorStoreIndex.from_documents(
                documents,
                storage_context=storage_context,
                embed_model=embedding_model,
                show_progress=False,
            )
        except Exception:
            for file_id in file_ids:
                chroma_collection.delete(where={"file_id": file_id})
            raise

    if app.state.index is None:
        embedding_model = get_embedding_model()
        vector_store = ChromaVectorStore(chroma_collection=chroma_collection)
        app.state.index = VectorStoreIndex.from_vector_store(
            vector_store,
            embed_model=embedding_model,
        )
    return app.state.index


@app.get("/")
@app.head("/")
def root():
    return {
        "status": "ok",
        "service": "LlamaIndex Knowledge Hub API",
        "docs": "/docs",
        "health": "/api/health",
    }


@app.get("/api/health")
def health_check():
    return {"status": "ok"}


@app.get("/api/files")
def list_files():
    files = [p.name for p in sorted(UPLOAD_DIR.glob("*")) if p.is_file() and p.suffix.lower() in ALLOWED_EXTS]
    return {"files": files}


@app.delete("/api/files/{filename}")
def delete_file(filename: str):
    if Path(filename).name != filename or Path(filename).suffix.lower() not in ALLOWED_EXTS:
        raise HTTPException(status_code=400, detail="Invalid file name.")

    target_path = UPLOAD_DIR / filename
    chroma_collection = get_chroma_collection()
    stored_ids = chroma_collection.get(where={"file_name": filename}).get("ids") or []
    if not target_path.exists() and not stored_ids:
        raise HTTPException(status_code=404, detail="File not found.")

    if stored_ids:
        chroma_collection.delete(ids=stored_ids)
    if target_path.exists():
        target_path.unlink()

    app.state.index = None
    app.state.chat_history.clear()
    return {"message": f"Removed {filename} and its indexed chunks."}


@app.post("/api/upload")
def upload_file(file: UploadFile = File(...)):
    if file.filename is None:
        raise HTTPException(status_code=400, detail="No file selected.")

    filename = Path(file.filename).name
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Allowed: PDF, TXT, DOCX, CSV, JSON, MD.",
        )

    target_path = UPLOAD_DIR / filename
    if target_path.exists():
        raise HTTPException(
            status_code=409,
            detail="A file with this name is already uploaded. Remove it before uploading a replacement.",
        )

    content = file.file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The selected file is empty.")

    file_id = uuid4().hex
    temp_path = UPLOAD_DIR / f".{file_id}{ext}"
    temp_path.write_bytes(content)

    try:
        index_uploaded_file(temp_path, filename, file_id)
        temp_path.replace(target_path)
    except ValueError as exc:
        chroma_collection = get_chroma_collection()
        chroma_collection.delete(where={"file_id": file_id})
        temp_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except RuntimeError as exc:
        chroma_collection = get_chroma_collection()
        chroma_collection.delete(where={"file_id": file_id})
        temp_path.unlink(missing_ok=True)
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except Exception as exc:
        chroma_collection = get_chroma_collection()
        chroma_collection.delete(where={"file_id": file_id})
        temp_path.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail=f"Index build failed: {exc}") from exc

    return {"message": f"Uploaded and indexed: {filename}", "filename": filename}


@app.get("/api/files/{filename}/content")
def get_file_content(filename: str):
    if Path(filename).name != filename or Path(filename).suffix.lower() not in ALLOWED_EXTS:
        raise HTTPException(status_code=400, detail="Invalid file name.")
    target_path = UPLOAD_DIR / filename
    if not target_path.exists():
        raise HTTPException(status_code=404, detail="File not found.")
    content = extract_text_from_file(target_path)
    return {"filename": filename, "content": content}


@app.post("/api/raw-text")
def save_and_index_raw_text(payload: dict):
    raw_content = str(payload.get("content", "")).strip()
    if not raw_content:
        raise HTTPException(status_code=400, detail="Content cannot be empty.")

    raw_title = str(payload.get("title", "")).strip()
    from datetime import datetime

    if not raw_title:
        raw_title = f"raw_data_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
    else:
        safe_name = Path(raw_title).name
        if not safe_name.lower().endswith(".txt"):
            safe_name = f"{safe_name}.txt"
        raw_title = safe_name

    target_path = UPLOAD_DIR / raw_title
    if target_path.exists():
        stem = target_path.stem
        raw_title = f"{stem}_{datetime.now().strftime('%H%M%S')}.txt"
        target_path = UPLOAD_DIR / raw_title

    file_id = uuid4().hex
    temp_path = UPLOAD_DIR / f".{file_id}.txt"
    temp_path.write_text(raw_content, encoding="utf-8")

    try:
        index_uploaded_file(temp_path, raw_title, file_id)
        temp_path.replace(target_path)
    except ValueError as exc:
        chroma_collection = get_chroma_collection()
        chroma_collection.delete(where={"file_id": file_id})
        temp_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except RuntimeError as exc:
        chroma_collection = get_chroma_collection()
        chroma_collection.delete(where={"file_id": file_id})
        temp_path.unlink(missing_ok=True)
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except Exception as exc:
        chroma_collection = get_chroma_collection()
        chroma_collection.delete(where={"file_id": file_id})
        temp_path.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail=f"Index build failed: {exc}") from exc

    return {"message": f"Saved and indexed raw text: {raw_title}", "filename": raw_title}


@app.post("/api/query")
def query_rag(payload: dict):
    question = str(payload.get("question", "")).strip()
    if not question:
        raise HTTPException(status_code=400, detail="Question cannot be empty.")

    try:
        index = ensure_index()
    except Exception as exc:
        raise HTTPException(status_code=503, detail=f"Could not load document embeddings: {exc}") from exc

    if index is None:
        return {
            "answer": "I don't have any information loaded yet. Please upload a document first.",
            "context": [],
            "sources": [],
        }

    retriever = index.as_retriever(similarity_top_k=4)
    nodes = retriever.retrieve(question)
    context_texts = [node.get_content() for node in nodes]

    if not context_texts:
        answer = "I don't have any information about this. Please ask something else."
        return {"answer": answer, "context": [], "sources": []}

    history = app.state.chat_history[-5:]
    history_text = "\n".join(f"Q: {h['question']}\nA: {h['answer']}" for h in history)
    system_prompt = (
        "You are a helpful assistant. Answer using only the provided context. "
        "If the answer is not found in the uploaded documents, reply exactly: "
        "I don't have any information about this. Please ask something else."
    )
    if history_text:
        system_prompt += f"\n\nConversation memory:\n{history_text}"

    try:
        llm = get_llm()
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    context = "\n\n---\n\n".join(context_texts)
    prompt = f"{system_prompt}\n\nContext:\n{context}\n\nQuestion: {question}"
    answer = str(llm.complete(prompt).text).strip()

    app.state.chat_history.append({"question": question, "answer": answer})
    sources = list(dict.fromkeys(node.metadata.get("file_name", "uploaded document") for node in nodes))
    return {"answer": answer, "context": context_texts, "sources": sources}


@app.get("/api/memory")
def get_memory():
    return {"memory": app.state.chat_history}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("app.main:app", host="0.0.0.0", port=8000, reload=True)
